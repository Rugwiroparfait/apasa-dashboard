"""Build and execute a source-preserving inspection notebook; never writes inputs."""
from pathlib import Path
import json, contextlib, io, traceback, sys
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
cells=[]
def md(s): cells.append({'cell_type':'markdown','metadata':{},'source':s.splitlines(True)})
def code(s): cells.append({'cell_type':'code','metadata':{},'source':s.splitlines(True),'execution_count':None,'outputs':[]})
md('''# APASA workbook inspection and candidate crosswalk
Review stage only — no cleaned observations, seasonal sums, conversions, or target comparisons.

Run cells in order. Inputs are opened read-only in intent and are never saved. SHA-256 checks verify unchanged bytes. Source labels, formulas, cached results, codes and blanks remain available in the cell audit. Candidate links are provisional interpretations, not verified equivalences.

The supplied directory has multiple NISR workbooks, not one consolidated history. Coverage is reported from evidence inside each sheet, with filename years identified separately. The register has six recurring indicator definitions and thirty seasonal records; stable local IDs are analytical identifiers, not official PSTA5 codes. Register baseline/target fields are kept in a separate register audit.
''')
code('''from pathlib import Path
import sys, json, csv, re, hashlib, html
from collections import Counter, defaultdict
import openpyxl
if Path('/tmp/apasa_workbook_deps').exists():
    sys.path.insert(0, '/tmp/apasa_workbook_deps')
import xlrd
ROOT = Path.cwd().resolve()
if not (ROOT / 'data').is_dir():
    ROOT = Path('/home/kali/Documents/Projects/apasa-dashboard')
OUT = ROOT / 'outputs/workbook_review_2026_10_28'
OUT.mkdir(parents=True, exist_ok=True)
inputs = sorted(p for p in (ROOT/'data').iterdir() if p.suffix.lower() in ('.xls','.xlsx','.xlsm'))
hash_file = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
before = {str(p.relative_to(ROOT)): hash_file(p) for p in inputs}
print('Excel inputs:', len(inputs))
for p in inputs: print(p.name)
print('Other data files (not treated as workbook observations):')
for p in sorted((ROOT/'data').iterdir()):
    if p not in inputs: print(p.name)
''')
md('''## 1. Inspect every sheet before interpreting relationships
`source_cells.json` is a sparse audit of every nonempty source cell, including cell references, formulas and cached values. Missing coordinates inside the reported used range remain blank; they are never zero-filled. Reported dimensions can include formatting-only columns. The inventory also records the actual nonempty extent, merged ranges, comments, notes, missing-value tokens and formula-cache problems. Previews are the first five nonempty rows, with coordinates; longer headers remain in the cell audit.
''')
code('''sheets=[]; inventory=[]; audits=[]
def dump_csv(name, rows):
    if not rows: return
    with (OUT/name).open('w', newline='', encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
def year_hits(text): return sorted(set(re.findall(r'(?<![0-9])(?:19|20)[0-9]{2}(?![0-9])',text)))
def refs(items): return ' | '.join(f"{c['cell']}={c['value']}" for c in items)
for p in inputs:
    if p.suffix.lower()=='.xls':
        wb=xlrd.open_workbook(p)
        source=[]
        for ws in wb.sheets():
            cs=[]
            for r in range(ws.nrows):
                for c in range(ws.ncols):
                    v=ws.cell_value(r,c)
                    if v != '': cs.append(dict(cell=f'{openpyxl.utils.get_column_letter(c+1)}{r+1}',row=r+1,col=c+1,value=v,cached_value=v,type=str(ws.cell_type(r,c)),comment=None))
            source.append((ws.name,ws.nrows,ws.ncols,cs,[str(x) for x in ws.merged_cells], 'Legacy XLS: xlrd returns stored results, not formula expressions; formula-cache completeness cannot be audited.'))
    else:
        wb=openpyxl.load_workbook(p,data_only=False,read_only=False)
        cached=openpyxl.load_workbook(p,data_only=True,read_only=False)
        source=[]
        for ws in wb:
            cs=[]
            for (r,c),obj in sorted(ws._cells.items()):
                if obj.value is not None:
                    cs.append(dict(cell=obj.coordinate,row=r,col=c,value=obj.value,cached_value=cached[ws.title].cell(r,c).value,type=obj.data_type,comment=obj.comment.text if obj.comment else None))
            source.append((ws.title,ws.max_row,ws.max_column,cs,[str(x) for x in ws.merged_cells.ranges],''))
    for name,nr,nc,cs,merged,limitation in source:
        byrow=defaultdict(list)
        for c in cs: byrow[c['row']].append(c)
        preview='\\n'.join(refs(byrow[r]) for r in sorted(byrow)[:5])
        top=[c for c in cs if c['row']<=8 and isinstance(c['value'],str)]
        # Titles often use formulas referencing a contents sheet: retain both formula and cached label.
        titles=[str(c['cached_value'] if c['type']=='f' and c['cached_value'] is not None else c['value']) for c in top if re.search(r'table|yield|production|summary|indicator|fertili|seed',str(c['cached_value'] or c['value']),re.I)]
        title=' | '.join(titles[:6])
        header=refs(top)
        notes=[c for c in cs if isinstance(c['value'],str) and re.search(r'(^\\s*(source|note|footnote)\\b|NISR|National Institute|\\*\\s*\\w)',c['value'],re.I)]
        units=[c for c in cs if isinstance(c['value'],str) and re.search(r'(kg\\s*/\\s*ha|mt/ha|ton|percentage|percent|\\(ha\\)|kg/|sqm|rwf|\\(%\\))',c['value'],re.I)]
        markers=Counter(str(c['value']) for c in cs if isinstance(c['value'],str) and c['value'].strip().lower() in ('-','--','…','...','n/a','na','nan','nil','..',':','(na)','(n/a)'))
        years=year_hits(header+' | '+title)
        filename_years=year_hits(p.name)
        seasons=sorted(set(re.findall(r'\\bSeason\\s*([ABC])\\b',header,re.I)))
        formula=[c for c in cs if c['type']=='f']
        missing_cache=sum(c['cached_value'] is None for c in formula)
        external=sum('[' in str(c['value']) for c in formula)
        row=dict(source_workbook=p.name,source_sheet_or_table=name,reported_rows=nr,reported_columns=nc,nonempty_rows=max((c['row'] for c in cs),default=0),nonempty_columns=max((c['col'] for c in cs),default=0),nonempty_cells=len(cs),title_or_measure_labels=title,header_years='; '.join(years),filename_years='; '.join(filename_years),season_labels='; '.join(seasons),unit_evidence=refs(units[:30]),geography_evidence=refs([c for c in top if re.search(r'district|province|rwanda|national|total',str(c['value']),re.I)]),group_and_denominator_evidence=refs([c for c in top if re.search(r'farmer|plot|land|SSF|LSF|crop|segment|site',str(c['value']),re.I)]),crop_label_evidence=refs([c for c in cs if isinstance(c['value'],str) and re.fullmatch(r'\s*(Maize|Paddy rice|Wheat|Cassava|Sweet potatoes?|Irish potatoes?|Beans|Bush beans?|Climbing beans?|Sorghum|Bananas|Cereals|Legumes.*|Tubers.*)\s*',c['value'],re.I)]),source_notes=refs(notes),missing_tokens=json.dumps(markers),formula_cells=len(formula),formula_without_cached_result=missing_cache,external_formula_cells=external,merged_ranges='; '.join(merged),reader_limitation=limitation,preview=preview)
        inventory.append(row)
        sheet=dict(workbook=p.name,sheet=name,cells=cs,header=header,title=title,years=years,filename_years=filename_years,seasons=seasons,inventory=row)
        sheets.append(sheet)
        audits.append(dict(source_workbook=p.name,sheet=name,reported_dimensions=[nr,nc],merged_ranges=merged,cells=cs,reader_limitation=limitation))
(OUT/'source_cells.json').write_text(json.dumps(audits,ensure_ascii=False,default=str),encoding='utf-8')
dump_csv('sheet_inventory.csv', inventory)
print('Sheets inspected:',len(inventory))
for p in inputs:
    ss=[s for s in inventory if s['source_workbook']==p.name]
    print(p.name, '—',len(ss),'sheets; header year labels:',sorted(set(y for s in ss for y in s['header_years'].split('; ') if y)))
print('Sheets with missing formula caches:',sum(s['formula_without_cached_result']>0 for s in inventory))
print('Sheets with external formulas:',sum(s['external_formula_cells']>0 for s in inventory))
''')
md('''## 2. Inspect register definitions and periods independently
All original baseline and target strings are retained in `register_fields_for_review.csv`. The register's own seasonal periods take precedence over assumptions about fiscal years. Its policy status and target provenance are unresolved: the supplied register is not an authenticated PSTA5 results framework. No external policy source has been used to validate it.
''')
code('''register=[]; indicators={}
for s in sheets:
    if 'PSTA5' not in s['workbook']: continue
    grid={(c['row'],c['col']):c['value'] for c in s['cells']}
    for r in range(2,s['inventory']['reported_rows']+1):
        name=grid.get((r,1))
        if not name: continue
        if name not in indicators:
            indicators[name]=dict(indicator_id=f'REG-{len(indicators)+1:02}',psta5_indicator=name,definition=grid.get((r,2)),unit=grid.get((r,3)),crop_or_group=grid.get((r,7)))
        register.append(dict(**indicators[name],register_sheet=s['sheet'],source_row=r,baseline_field=grid.get((1,4)),baseline_value_as_written=grid.get((r,4)),target_field=grid.get((1,5)),target_value_as_written=grid.get((r,5)),stated_source=grid.get((r,6)),policy_provenance='UNRESOLVED: supplied register field labels retained; not independently verified PSTA5 targets',geography='Not explicitly stated in register'))
dump_csv('register_fields_for_review.csv',register)
for x in indicators.values(): print(x)
print('Register records:',len(register))
''')
md('''## 3. Candidate table/year relationships
Candidate inclusion is deliberately broad: sheets containing production, yield/productivity, harvested/cultivated area, improved seeds, inorganic fertilizer, or summary main indicators in their initial labels. Each of the six distinct register definitions is considered against every included table/year; unrelated topic pairs are contextual. Other sheets remain visible in the inventory and exclusion audit. A year mentioned only in a filename is provisional, and a year in a multiyear header does not establish the reference period of every cell. No observation values are extracted into this crosswalk.

POSSIBLE MATCH requires direct concept evidence and, for input use, explicit farmer denominator and requested farmer type. Plot/land measures and seed/fertilizer composition or sources are context. Missing denominator/group evidence is UNRESOLVED. Yield labels remain provisional pending crop and harvested-area denominator verification. Mixed summary tables are UNRESOLVED until their individual blocks are mapped. The exact source strings and coordinate evidence accompany every row; analytical topic labels do not replace source categories.
''')
code('''crosswalk=[]; excluded=[]
for s in sheets:
    if 'PSTA5' in s['workbook']: continue
    h=(s['header']+' | '+s['title']).lower(); title=s['title'].lower()
    families=[]
    if re.search(r'production|produced',h): families.append('production')
    if re.search(r'yield|productivity',h): families.append('yield')
    if re.search(r'harvested area|cultivated area',h): families.append('area')
    if re.search(r'improved seed|type of seeds|types of seeds|seed use|use of seeds|seeds used',h): families.append('seed')
    if re.search(r'inorganic fertili',h): families.append('fertilizer')
    summary=bool(re.search(r'summary.*main indicator|main sas indicator|sas main indicator',h))
    if summary: families.append('summary')
    if not families or re.search(r'list of table',s['sheet'],re.I) or s['sheet'].lower()=='summary':
        excluded.append(dict(source_workbook=s['workbook'],source_sheet=s['sheet'],reason='Contents/navigation or no candidate concept in first eight rows; inspect inventory for scope extension',exact_header=s['header']))
        continue
    yearlist=s['years'] or s['filename_years'] or ['UNRESOLVED']
    for ind in indicators.values():
        name=ind['psta5_indicator']; target_topic='production' if name.startswith('Quantity') else 'yield' if name.startswith('Productivity') else 'seed' if 'seeds' in name else 'fertilizer'
        group='SSF' if name.startswith('Small') else 'LSF' if name.startswith('Large') else ''
        status='CONTEXT ONLY'; reason='Different measured concept; potential background or supporting input only.'
        if summary:
            status='UNRESOLVED'; reason='Mixed main-indicator summary: map exact row, year/season columns, unit and population before determining relationship.'
        elif target_topic=='production' and re.search(r'use of (crop )?production|production use|share of crop production|sampling error',title):
            reason='Production allocation/use shares or sampling-error statistics do not directly measure crop production volume.'
        elif target_topic in families:
            if target_topic in ('production','yield'):
                status='POSSIBLE MATCH'; reason='Direct production/yield concept, subject to exact crop, geography, population, unit and period checks.'
            else:
                subtype=bool(re.search(r'source of|sources of|type of inorganic|types of inorganic|source.*seed',title))
                farmer=bool(re.search(r'percentage of farmers|farmers who|farmer type|farmers.*us|farmers.*appl',h))
                hasgroup=bool(re.search(r'\\b'+group+r'\\b',h,re.I) or re.search('small.?scale' if group=='SSF' else 'large.?scale',h,re.I))
                if subtype:
                    reason='Source/composition of input use does not measure farmer adoption prevalence.'
                elif farmer and hasgroup:
                    status='POSSIBLE MATCH'; reason='Farmer-use concept and requested group appear in headers; select farmer block only, never plot or land blocks.'
                elif 'plot' in title and not farmer:
                    reason='Plot denominator differs from the register farmer denominator.'
                else:
                    status='UNRESOLVED'; reason='Input-use topic present, but farmer denominator and/or requested farmer group is not explicit enough in table header.'
        elif target_topic=='yield' and ('production' in families or 'area' in families):
            reason='Supporting numerator/area context only; do not derive yield without matched crop, harvested-area denominator and documented estimator.'
        for year in yearlist:
            note=reason+' Period alignment with each register season must be checked; no annual aggregation. Register policy target provenance remains unresolved.'
            if not s['years']:
                note+=' Year inferred from filename only, not verified in header.'
                if status=='POSSIBLE MATCH': status='UNRESOLVED'
            if s['inventory']['formula_without_cached_result']:
                note+=' Missing formula caches: do not treat blank cached results as observed values.'
            conversion='No conversion applied; verify units in exact block.'
            if target_topic=='yield': conversion='If source is kg/ha and register MT/ha, divide by 1000 only after denominator/crop verification; no conversion applied.'
            crosswalk.append(dict(indicator_id=ind['indicator_id'],psta5_indicator=name,psta5_definition=ind['definition'],psta5_unit=ind['unit'],psta5_crop_or_group=ind['crop_or_group'],register_records='; '.join(r['register_sheet']+' row '+str(r['source_row']) for r in register if r['indicator_id']==ind['indicator_id']),source_workbook=s['workbook'],source_sheet_or_table=s['sheet'],year=year,year_basis='header evidence (may be comparative; verify column)' if s['years'] else 'filename only — unresolved',season='; '.join(s['seasons']) or 'UNRESOLVED',nisr_source_label=s['title'],exact_header_evidence=s['header'],definition_and_calculation='Not established from table label; inspect source cells/formulas and survey methodology. No calculation performed.',unit_evidence=s['inventory']['unit_evidence'] or 'Not explicit in inspected labels',crop_or_group_evidence=s['inventory']['group_and_denominator_evidence'] or 'UNRESOLVED; see source cells',exact_crop_label_evidence=s['inventory']['crop_label_evidence'],geography_evidence=s['inventory']['geography_evidence'] or 'UNRESOLVED',denominator='Keep farmer, plot, land and harvested-area denominators separate; exact block verification required',source_note=s['inventory']['source_notes'] or 'No explicit source note detected; local workbook provenance only',missing_or_not_applicable=s['inventory']['missing_tokens']+'; blanks preserved; token meanings unverified',unit_conversion_required=conversion,value_role='NISR table candidate; no observed value extracted; register baselines/targets are separate',match_status=status,comparability_note=note))
dump_csv('candidate_crosswalk.csv',crosswalk)
dump_csv('noncandidate_sheets.csv',excluded)
print('Candidate rows:',len(crosswalk)); print(dict(Counter(r['match_status'] for r in crosswalk)))
for ind in indicators.values():
    rr=[r for r in crosswalk if r['indicator_id']==ind['indicator_id']]
    print(ind['indicator_id'],ind['psta5_indicator'],dict(Counter(r['match_status'] for r in rr)))
assert len({(r['indicator_id'],r['source_workbook'],r['source_sheet_or_table'],r['year']) for r in crosswalk})==len(crosswalk)
assert {r['match_status'] for r in crosswalk} <= {'POSSIBLE MATCH','CONTEXT ONLY','UNRESOLVED'}
assert set(ind['indicator_id'] for ind in indicators.values()) == set(r['indicator_id'] for r in crosswalk)
''')
md('''## 4. Readable review artifacts and preservation checks
The HTML report displays the entire inventory and candidate crosswalk, with search filtering and source previews. No verified tidy observation CSV is created at this review gate. CSVs here contain metadata, source register fields and candidate relationships only.
''')
code('''def table(rows, fields):
    return '<table><thead><tr>'+''.join('<th>'+html.escape(k)+'</th>' for k in fields)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(str(r.get(k,'')))+'</td>' for k in fields)+'</tr>' for r in rows)+'</tbody></table>'
style='<style>body{font:15px system-ui;margin:30px;color:#20332d}table{border-collapse:collapse;width:100%;margin:20px 0}td,th{border:1px solid #ccc;padding:8px;text-align:left;vertical-align:top;white-space:pre-wrap;min-width:100px}th{background:#e6eee8}section{overflow:auto}input{padding:10px;width:65%}h2{margin-top:35px}</style>'
intro=f'<h1>APASA workbook review — provisional</h1><p>{len(inputs)} Excel files; {len(inventory)} sheets; {len(indicators)} distinct register definitions; {len(register)} seasonal register records; {len(crosswalk)} candidate table/year relationships.</p><p>Review gate: no cleaned observations exported. POSSIBLE MATCH is not verified comparability. Register target provenance remains unresolved. Years in comparative headers require column-level review.</p><p>Search each table by indicator, workbook, sheet or year. Full coordinate evidence is in the CSV and source_cells.json.</p>'
report=style+intro+'<p><a href="review_note.md">Read the unresolved issues and source-cell examples</a></p>'
for label,rows,fields in [
    ('Inventory',inventory,['source_workbook','source_sheet_or_table','reported_rows','reported_columns','nonempty_rows','nonempty_columns','title_or_measure_labels','header_years','filename_years','season_labels','preview','source_notes','missing_tokens','formula_without_cached_result','external_formula_cells','reader_limitation']),
    ('Register definitions and separately retained baseline/target fields',register,list(register[0])),
    ('Candidate crosswalk',crosswalk,['indicator_id','psta5_indicator','source_workbook','source_sheet_or_table','year','year_basis','season','nisr_source_label','match_status','comparability_note','unit_conversion_required','crop_or_group_evidence'])]:
    report+='<h2>'+label+'</h2><input placeholder="Filter this table"><section>'+table(rows,fields)+'</section>'
report+="<script>document.querySelectorAll('input').forEach(el=>el.addEventListener('input',()=>{for(const r of el.nextElementSibling.querySelectorAll('tbody tr'))r.hidden=!r.textContent.toLowerCase().includes(el.value.toLowerCase())}));</script>"
(OUT/'review.html').write_text('<!doctype html><meta charset="utf-8"><title>APASA workbook review</title>'+report,encoding='utf-8')
after={str(p.relative_to(ROOT)):hash_file(p) for p in inputs}
assert before==after, 'Input changed during inspection'
(OUT/'input_manifest.json').write_text(json.dumps({'before':before,'after':after,'unchanged':before==after},indent=2))
print('Verified all input SHA-256 hashes unchanged. Open review.html for the full inventory and crosswalk.')
''')
md((OUT/'review_note.md').read_text())
ns={'__name__':'__main__'}
for i,c in enumerate(cells):
 if c['cell_type']=='code':
  output=io.StringIO()
  with contextlib.redirect_stdout(output): exec(compile(''.join(c['source']),f'cell_{i}','exec'),ns)
  c['execution_count']=sum(x['cell_type']=='code' for x in cells[:i+1])
  c['outputs']=[{'output_type':'stream','name':'stdout','text':output.getvalue().splitlines(True)}]
  print(output.getvalue(),end='')
notebook={'cells':cells,'metadata':{'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'},'language_info':{'name':'python','version':sys.version.split()[0]}},'nbformat':4,'nbformat_minor':5}
for i,c in enumerate(cells): c['id']=f'apasa-review-{i:02}'
(OUT/'apasa_workbook_review.ipynb').write_text(json.dumps(notebook,indent=1,ensure_ascii=False),encoding='utf-8')
